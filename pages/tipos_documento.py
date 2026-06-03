import streamlit as st
import pandas as pd
import io
import csv
from openpyxl import load_workbook
from streamlit_echarts import st_echarts

st.set_page_config(page_title="Tipos de Documento", layout="wide")

@st.cache_data
def load_data():
    wb = load_workbook("scopus_g2_ventas.xlsx", read_only=True)
    ws = wb.active
    lines = []
    for row in ws.iter_rows(values_only=True):
        cells = [str(c) for c in row if c is not None]
        if cells:
            lines.append(",".join(cells))
    csv_content = "\n".join(lines[1:])
    reader = csv.DictReader(io.StringIO(csv_content))
    rows = list(reader)
    df = pd.DataFrame(rows)
    df.columns = [c.strip('"').strip() for c in df.columns]
    return df

df = load_data()

st.page_link("pages/dashboard.py", label="⬅️ Volver al Dashboard Principal", icon="🏠")
st.title("📄 Distribución por Tipo de Documento")
st.markdown("---")

col1, col2 = st.columns([2, 1])

with col1:
    tipo_counts = df["Document Type"].value_counts()
    tipo_data = [{"value": int(v), "name": k} for k, v in tipo_counts.items()]

    option_tipo = {
        "tooltip": {"trigger": "item", "formatter": "{b}: {c} ({d}%)"},
        "color": ["#1D9E75", "#378ADD", "#D85A30", "#BA7517", "#533AB7"],
        "legend": {"orient": "vertical", "right": "5%", "top": "center"},
        "series": [{
            "type": "pie",
            "radius": ["40%", "70%"],
            "center": ["40%", "50%"],
            "data": tipo_data,
            "label": {"show": True, "formatter": "{b}\n{d}%"},
            "itemStyle": {"borderRadius": 5, "borderColor": '#fff', "borderWidth": 2}
        }]
    }
    st_echarts(option_tipo, height="500px")

with col2:
    st.subheader("💡 Interpretación")
    st.write("""
    Este gráfico de torta muestra que tipo de formato eligen los investigadores para publicar sus articulos, nos permite visualizar qué tipos de documentos (artículos, revisiones, ponencias de congresos) son más frecuentes en el dataset.

    **Puntos clave/Insights:**
    * **Innovación acelerada: La mayoría de las investigaciones se presentan en conferencias esto porque en el campo de la Ciencia de Datos e Inteligencia Artificial, las tecnologías evolucionan tan rápido que los autores prefieren presentar sus resultados inmediatamente en congresos internacionales.
    * **Los Artículos de Revista Científica:** Publicar en una revista científica tradicional (lo que se clasifica estrictamente como Article) es un proceso muy largo y riguroso. Puede tomar entre 1 y 2 años desde que se envía el texto hasta que se publica. Para que un modelo predictivo sea aceptado aquí, tiene que estar increíblemente pulido, probado en muchísimos escenarios distintos
    """)
