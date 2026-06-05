import streamlit as st
import pandas as pd
import io
import csv
from openpyxl import load_workbook
from streamlit_echarts import st_echarts

st.set_page_config(page_title="Fuentes Científicas", layout="wide")

@st.cache_data
def load_data():
    wb = load_workbook("scopus_g2_ventas.xlsx", read_only=True)
    ws = wb.active
    lines = []
    for row in ws.iter_rows(values_only=True):
        cells = [str(c) for c in row if c is not None]
        if cells:
            lines.append(",".join(cells))
    csv_content = "\n".join(lines[1:])
    reader = csv.DictReader(io.StringIO(csv_content))
    rows = list(reader)
    df = pd.DataFrame(rows)
    df.columns = [c.strip('"').strip() for c in df.columns]
    return df

df = load_data()

st.page_link("pages/dashboard.py", label="⬅️ Volver al Dashboard Principal", icon="🏠")
st.title("Fuentes científicas más frecuentes")
st.markdown("---")

col1, col2 = st.columns([2, 1])

with col1:
    top_fuentes = df["Source title"].value_counts().head(10)
    fuentes_labels = [str(f)[:60] + "…" if len(str(f)) > 60 else str(f) for f in top_fuentes.index]
    fuentes_vals = top_fuentes.values.tolist()
    fuente_data = list(zip(fuentes_labels[::-1], fuentes_vals[::-1]))

    option_fuentes = {
        "tooltip": {"trigger": "axis", "formatter": "{b}: {c} publicaciones"},
        "xAxis": {"type": "value", "minInterval": 1},
        "yAxis": {"type": "category", "data": [f[0] for f in fuente_data], "axisLabel": {"width": 300, "overflow": "truncate"}},
        "series": [{
            "type": "bar",
            "data": [f[1] for f in fuente_data],
            "itemStyle": {"color": "#0F6E56", "borderRadius": [0, 4, 4, 0]},
            "label": {"show": True, "position": "right"}
        }],
        "grid": {"left": "2%", "right": "10%", "bottom": "5%", "containLabel": True}
    }
    st_echarts(option_fuentes, height="600px")

with col2:
    st.subheader("Interpretación")
    st.write("""  
Este gráfico nos muestra la distribucion de las fuentes cientificas es decir muestra en qué revistas científicas, congresos o conferencias académicas se están publicando los artículos. Al observar el grafico podemos concluir:

**Puntos clave/Insights:**
* **Dominio netamente tecnológico:** La investigación no se está publicando en revistas tradicionales de marketing o ventas, sino que está dominada casi por completo por las ciencias de la computación y las ingenierías, esto confirma que mejorar la precisión de las ventas es hoy, fundamentalmente, un desafío de arquitectura de datos e inteligencia artificial.
* **Concentración en el liderazgo:** Existe una clara preferencia por dos medios principales para publicar sobre este tema: *Communications in Computer and Information Science* y *Lecture Notes in Computer Science*, ambas destacándose del resto con 3 publicaciones cada una. 
* **Velocidad de innovación:** La gran presencia de actas de conferencias (*Proceedings*, *Conferences*) indica que este es un campo de estudio muy dinámico. En este caso la minería de datos y la IA, los investigadores prefieren compartir y presentar sus nuevos algoritmos predictivos en conferencias para compartir los resultados más rápido.
""")
