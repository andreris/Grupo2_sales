import streamlit as st
import pandas as pd
import io
import csv
from openpyxl import load_workbook
from streamlit_echarts import st_echarts

st.set_page_config(page_title="Top Citados", layout="wide")

@st.cache_data
def load_data():
    wb = load_workbook("scopus_g2_ventas.xlsx", read_only=True)
    ws = wb.active
    lines = []
    for row in ws.iter_rows(values_only=True):
        cells = [str(c) for c in row if c is not None]
        if cells:
            lines.append(",".join(cells))
    csv_content = "\n".join(lines[1:])
    reader = csv.DictReader(io.StringIO(csv_content))
    rows = list(reader)
    df = pd.DataFrame(rows)
    df.columns = [c.strip('"').strip() for c in df.columns]
    df["Cited by"] = pd.to_numeric(df["Cited by"], errors="coerce").fillna(0)
    return df

df = load_data()

st.page_link("pages/dashboard.py", label="⬅️ Volver al Dashboard Principal", icon="🏠")
st.title("🏆 Top 10 Artículos Más Citados")
st.markdown("---")

col1, col2 = st.columns([2, 1])

with col1:
    top_citados = df.nlargest(10, "Cited by")[["Title", "Cited by"]].copy()
    top_citados["Titulo corto"] = top_citados["Title"].str.strip('"').str[:60] + "…"
    top_citados = top_citados.sort_values("Cited by", ascending=True)

    option_citas = {
        "tooltip": {"trigger": "axis", "formatter": "{b}: {c} citas"},
        "xAxis": {"type": "value"},
        "yAxis": {"type": "category", "data": top_citados["Titulo corto"].tolist(), "axisLabel": {"width": 300, "overflow": "truncate"}},
        "series": [{
            "type": "bar",
            "data": top_citados["Cited by"].astype(int).tolist(),
            "itemStyle": {"color": "#533AB7", "borderRadius": [0, 4, 4, 0]},
            "label": {"show": True, "position": "right"}
        }],
        "grid": {"left": "2%", "right": "10%", "bottom": "5%", "containLabel": True}
    }
    st_echarts(option_citas, height="600px")

with col2:
    st.subheader("💡 Interpretación")
    st.write("""En este gráfico podemos ver cuáles son los artículos de nuestro dataset que más veces han sido citados. Básicamente, nos ayuda a entender qué investigaciones son las más importantes o tomadas como referencia en el tema de predicción de ventas.

**Puntos clave/Insights:**
* **Los artículos más influyentes:** Vemos que hay un artículo que resalta muchísimo sobre el resto. Esto normalmente pasa porque propone un modelo o una metodología que se vuelve la base para que otros hagan sus propios experimentos.
* **Por qué importan las citas:** En la investigación, si te citan mucho es porque tu trabajo sirve. Que estos artículos tengan tantas menciones nos dice que sus modelos de machine learning realmente dan buenos resultados y la comunidad confía en ellos.
* **Un estudio clave:** La gran diferencia de citas entre el primer lugar y los demás nos demuestra que ese estudio en particular marcó un antes y un después, y es casi una lectura obligatoria si queremos entender cómo predecir ventas con IA.
""")
