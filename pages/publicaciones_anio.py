import streamlit as st
import pandas as pd
import io
import csv
from openpyxl import load_workbook
from streamlit_echarts import st_echarts

st.set_page_config(page_title="Publicaciones por año", layout="wide")

@st.cache_data
def load_data():
    wb = load_workbook("scopus_g2_ventas.xlsx", read_only=True)
    ws = wb.active
    lines = []
    for row in ws.iter_rows(values_only=True):
        cells = [str(c) for c in row if c is not None]
        if cells:
            lines.append(",".join(cells))
    csv_content = "\n".join(lines[1:])
    reader = csv.DictReader(io.StringIO(csv_content))
    rows = list(reader)
    df = pd.DataFrame(rows)
    df.columns = [c.strip('"').strip() for c in df.columns]
    df["Year"] = pd.to_numeric(df["Year"], errors="coerce")
    df = df.dropna(subset=["Year"])
    df["Year"] = df["Year"].astype(int)
    return df

df = load_data()

st.page_link("pages/dashboard.py", label="⬅️ Volver al Dashboard Principal", icon="🏠")
st.title("Análisis de Publicaciones por Año")
st.markdown("---")

col1, col2 = st.columns([2, 1])

with col1:
    pub_por_anio = df["Year"].value_counts().sort_index()
    years_list = [str(y) for y in pub_por_anio.index.tolist()]
    counts_list = pub_por_anio.values.tolist()
    bar_colors = ["#FF5F00" if int(y) >= 2020 else "#1847A5" for y in years_list] 

    option_anio = {
        "tooltip": {"trigger": "axis", "formatter": "{b}: {c} publicaciones"},
        "xAxis": {"type": "category", "data": years_list, "axisLabel": {"rotate": 45}},
        "yAxis": {"type": "value", "minInterval": 1},
        "series": [{
            "type": "bar",
            "data": [{"value": v, "itemStyle": {"color": c}} for v, c in zip(counts_list, bar_colors)],
            "itemStyle": {"borderRadius": [4, 4, 0, 0]},
            "label": {"show": True, "position": "top"}
        }],
        "grid": {"left": "5%", "right": "5%", "bottom": "15%", "containLabel": True}
    }
    st_echarts(option_anio, height="500px")

with col2:
    st.subheader("Interpretación")
    st.write("""
    Este gráfico muestra la cantidad de artículos científicos publicados por año relacionados con la predicción de ventas mediante inteligencia artificial, inteligencia de negocios y minería de datos. Permite identificar los años con mayor producción científica sobre el tema.
    
    **Puntos clave/Insights:**
    * **Tendencia:** Se observa un crecimiento en los últimos años, en especial desde el año 2020 del uso y conocimiento sobre la ia, lo que demuestra que es un área que esta en crecimiento.
    * **Tecnología :** El aumento del uso de machine learning y la adopción de la ia en las empresas se ve claro, más aún en el año 2025.
    * **Larga espera :** El tiempo que requeria el podeer recolectar datos, escribir los papers y pasar por los rigurosas revisiones, hicieron que el 2025 se convirtió en el año con mayores publicaciones, esto con todo lo investigado durante el 2023 y 2024 principalmente.
    """)